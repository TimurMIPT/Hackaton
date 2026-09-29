from pathlib import Path
from datetime import datetime, date
import argparse
import csv
import hashlib
import json
import openpyxl

parser = argparse.ArgumentParser(description='Пересобрать rules.json и work_mapping.csv из Excel-перечня.')
parser.add_argument('--source', type=Path, default=Path.cwd() / 'Сводный перечень строительных работ_ЛТЦ.xlsx',
                    help='Путь к исходному Excel-файлу.')
parser.add_argument('--out', type=Path, default=Path.cwd(),
                    help='Каталог для rules.json и work_mapping.csv.')
args = parser.parse_args()
SRC = args.source.resolve()
OUT = args.out.resolve()

machines = {
 'dump_truck': {'name':'Самосвал','detector':True,'aliases':['самосвал','самосвалы','самосвала','dump truck','dump_truck']},
 'bulldozer': {'name':'Бульдозер','detector':True,'aliases':['бульдозер','бульдозеры','бульдозера','bulldozer']},
 'excavator': {'name':'Экскаватор','detector':True,'aliases':['экскаватор','экскаваторы','экскаватора','экскавтор','excavator']},
 'mobile_crane': {'name':'Автокран','detector':True,'aliases':['автокран','автокраны','автокрана','автомобильный кран','mobile crane','mobile_crane','truck crane']},
 'mixer': {'name':'Автобетоносмеситель','detector':True,'aliases':['автобетоносмеситель','автобетоносмесители','автобетоносмесителя','автобетономешалка','автобетономешалки','автобетономешалок','бетономешалка','бетономешалки','бетономешалок','миксер','mixer','concrete mixer']},
 'tower_crane': {'name':'Кран','detector':True,'aliases':['кран','краны','крана','башенный кран','башенные краны','tower_crane','tower crane','crane']},
 'roller': {'name':'Каток','detector':True,'aliases':['каток','катки','катка','катков','roller','road roller']},
 'pile_rig': {'name':'Свайная установка','detector':True,'aliases':['свайная установка','свайные установки','свайных установок','свайной установки','буровая установка','сваебойная установка','pile rig','pile_rig']},
}
classes = {
 'C01': {'name':'Подготовка территории и демонтаж','profiles':['E','L'],'observability':'Частично','note':'Техника похожа на земляные работы. Геодезия, переселение и оснащение площадки вынесены в C10.'},
 'C02': {'name':'Земляные работы и основания','profiles':['E','R'],'observability':'Частично','note':'Котлован, выемка, планировка, засыпка и подготовка основания. Направление перемещения грунта по списку техники неизвестно.'},
 'C03': {'name':'Сваи и крепление грунта','profiles':['P','L','B'],'observability':'Частично','note':'Свайная установка распознаётся моделью. Шпунт и закрепление грунта требуют уточнения технологии.'},
 'C04': {'name':'Фундамент и подземные конструкции','profiles':['B','L'],'observability':'Частично','note':'Бетонирование и монтаж наблюдаемы косвенно. Армирование, опалубка и гидроизоляция могут выполняться без видимой тяжёлой техники.'},
 'C05': {'name':'Каркас и надземные конструкции','profiles':['B','L'],'observability':'Частично','note':'Бетонирование и монтаж. Профиль техники совпадает с C04; высота, геометрия здания и план нужны для их разделения.'},
 'C06': {'name':'Наружные инженерные сети','profiles':['E','L','B'],'observability':'Слабо','note':'Земляные и монтажные работы при открытой прокладке. По составу техники неотличимы от других земляных работ.'},
 'C07': {'name':'Отделка, кровля и внутренние системы','profiles':[],'observability':'Не определяется','note':'Черновая и чистовая отделка объединены. Также кладка, изоляция, фасадная отделка и внутренние сети. Нужны признаки конструкций, материалов или рабочих.'},
 'C08': {'name':'Дороги и благоустройство','profiles':['R','E','B','L'],'observability':'Частично','note':'Механизированные подоперации проверяются отдельно. Озеленение, разметка и малые ручные работы не требуют видимой тяжёлой техники.'},
 'C09': {'name':'Тоннели и специальные путевые работы','profiles':[],'observability':'Не определяется','note':'ТПМК, тоннельная проходка и железнодорожный путь. Заданных типов недостаточно; путевой бетон отмечен отдельно как бетонная подоперация.'},
 'C10': {'name':'Организация, поставки и оснащение','profiles':[],'observability':'Не определяется','note':'Переселение, геодезия, заказ и поставка оборудования, видеонаблюдение и СКУД. Нужны документы и другие наблюдения.'},
}
profiles = {
 'E': {'name':'Земляные работы','anchors':['excavator','bulldozer'],'weights':{'dump_truck':.65,'bulldozer':1.,'excavator':1.,'mobile_crane':.05,'roller':.15},'candidate_classes':['C01','C02','C06','C08']},
 'R': {'name':'Уплотнение и дорожные основания','anchors':['roller'],'weights':{'dump_truck':.35,'bulldozer':.45,'excavator':.15,'roller':1.},'candidate_classes':['C02','C08']},
 'P': {'name':'Свайные и буровые работы','anchors':['pile_rig'],'weights':{'dump_truck':.1,'bulldozer':.05,'excavator':.1,'mobile_crane':.25,'mixer':.35,'tower_crane':.1,'pile_rig':1.},'candidate_classes':['C03']},
 'B': {'name':'Бетонирование','anchors':['mixer'],'weights':{'dump_truck':.05,'mobile_crane':.15,'mixer':1.,'tower_crane':.25,'pile_rig':.15},'candidate_classes':['C03','C04','C05','C06','C08','C09']},
 'L': {'name':'Подъём и монтаж','anchors':['mobile_crane','tower_crane'],'weights':{'dump_truck':.05,'mobile_crane':1.,'mixer':.15,'tower_crane':1.,'pile_rig':.05},'candidate_classes':['C01','C03','C04','C05','C06','C08']},
}

def req(any_of, note='', haul=False):
    return {'any_of':any_of,'min_count':1,'note':note,'haul_only':haul}

rules = {
 'none': {'name':'Техника не даёт проверяемого признака','profiles':[],'required':[],'note':'Отсутствие тяжёлой техники не подтверждает и не опровергает выполнение работы.'},
 'earth_haul': {'name':'Выемка с вывозом','profiles':['E'],'required':[req(['excavator']),req(['dump_truck'],'Только сценарий с вывозом грунта или мусора.',True)],'note':'Принят сценарий механизированной выемки с вывозом. При складировании на площадке используйте --no-haul.'},
 'earth': {'name':'Перемещение и планировка грунта','profiles':['E'],'required':[req(['excavator','bulldozer'],'Достаточен один из этих типов.')],'note':'Самосвал допускается, но обязателен только в сценарии с вывозом.'},
 'base': {'name':'Устройство основания или засыпка','profiles':['E','R'],'required':[],'note':'Выемка, перемещение и уплотнение сменяют друг друга; одновременно все машины не требуются.'},
 'compact': {'name':'Уплотнение катком','profiles':['R'],'required':[req(['roller'],'Для принятого сценария с катком.')],'note':'На малых площадях возможны виброплиты и трамбовки вне перечня детектора; правило нужно изменить под проект.'},
 'pile': {'name':'Бурение или устройство свай','profiles':['P'],'required':[req(['pile_rig'])],'note':'Миксер условен: зависит от вида свай и подоперации. Свайная установка распознаётся моделью.'},
 'pile_support': {'name':'Крепление котлована и грунта','profiles':['P','L','B'],'required':[],'note':'Бурение, шпунт, вибропогружение и инъектирование требуют разной техники. Это допускаемые признаки, а не обязательная одновременная комбинация.'},
 'concrete': {'name':'Бетонирование с привозным бетоном','profiles':['B'],'required':[req(['mixer'],'Для сценария поставки готовой бетонной смеси.')],'note':'Миксер может отсутствовать между доставками. Бетононасос не распознаётся и не проверяется. Для другого снабжения измените правило.'},
 'concrete_cycle': {'name':'Цикл монолитных работ','profiles':['B','L'],'required':[],'note':'Опалубка, армирование и бетонирование выполняются последовательно. На произвольном снимке миксер не обязателен.'},
 'lift': {'name':'Механизированный подъём и монтаж','profiles':['L'],'required':[req(['mobile_crane','tower_crane'],'Достаточен автокран или кран.')],'note':'Для монтажа, предполагающего тяжёлый подъём. Размеры элементов и технология должны соответствовать этому сценарию.'},
 'lift_optional': {'name':'Армирование и подготовка к монтажу','profiles':['L'],'required':[],'note':'Кран может участвовать в подаче или установке элементов, но вязка и сборка выполняются и без него. Отсутствие крана не считается отклонением.'},
 'network': {'name':'Наружные сети, открытая прокладка','profiles':['E','L'],'required':[],'note':'Подоперация может быть ручной; траншея, трубы и местоположение информативнее одного состава машин.'},
 'road': {'name':'Дороги и основания, общий цикл','profiles':['E','R','B','L'],'required':[],'note':'Выберите конкретную подоперацию для проверки наличия техники.'},
}

wb = openpyxl.load_workbook(SRC,read_only=True,data_only=True)
ws = wb.worksheets[0]
object_types = [str(ws.cell(3,c).value) for c in range(3,12)]
raw = []
for row in range(4,381):
    cells = list(next(ws.iter_rows(min_row=row,max_row=row,values_only=True)))
    if not cells[1]:
        continue
    value = cells[0]
    repaired = isinstance(value,(datetime,date))
    code = f'{value.day}.{value.month}' if repaired else str(value or '').strip().rstrip('.')
    raw.append({'work_id':f'R{row:03d}','source_row':row,'source_code':code,'source_code_raw':str(value or ''),'code_repaired':repaired,'name':str(cells[1]).strip(),'object_marks':list(cells[2:11])})

# Explicit, reviewed row ranges refer to this exact source file, not to text heuristics.
mapping = {}
def put(rows, cid, rid, comment=''):
    for row in rows:
        mapping[row] = [cid,rid,comment]

put(range(4,43),'C10','none')
put([4,19,20,21,22,24,25,27,28,31],'C01','earth')
put([19,20,25],'C01','earth_haul')
put([22,28],'C01','none','Технология может не требовать тяжёлой техники.')
put([21],'C01','none','Вырубка может выполняться ручным инструментом; заданная тяжёлая техника не обязательна.')
put([24,31],'C01','lift')
put(range(6,19),'C06','network','Вынос наружных сетей до основных работ.')
put([29],'C06','network')
put([30],'C08','road','Временная дорога на подготовительном этапе, не признак позднего благоустройства.')
put([43,46,242],'G00','none','Сводный раздел включает разные классы. Нужна конкретная работа.')
put([44,45],'C09','none','ТПМК не относится к заданным типам техники.')
put(range(47,96),'C04','concrete_cycle')
put([47,65,72,73],'C02','earth_haul')
put([49,52,53,64,66,67,76,77,78],'C02','base')
put([67],'C02','earth')
put([68],'C02','compact')
put([48,51,60,74,75],'C03','pile')
put([59,61,63,70,71],'C03','pile_support')
put([58,69],'C04','lift','Кран — косвенный признак доставки или установки. Вязка арматуры вручную не наблюдается.')
put([58],'C04','lift_optional')
put(range(79,83),'C08','road')
put([81],'C08','earth')
put([82],'C08','compact')
put([54],'C04','concrete')
put(range(91,96),'C04','none','Гидроизоляция и инъектирование не определяются заданным составом техники.')
put(range(96,242),'C07','none')
put([96,100,102,104,106,167,172,173,174,175,176,181,182],'C05','concrete_cycle')
put([99,101,107,109,110,119,183,184,185,189,205,206,216,241],'C05','lift')
put([129,130,108],'C05','concrete')
put([101,107],'C05','lift_optional')
put([97,98],'C10','none')
put(range(136,171),'C08','none')
put([136,137,141,148,149,150],'C08','road')
put([138],'C08','earth')
put([139],'C08','compact')
put([142,143],'C08','concrete_cycle')
put([144,145,146,147,154,155,166,169],'C08','lift')
put([167],'C05','concrete_cycle')
put([168,170],'C07','none')
put([171],'C08','base')
put([156,157,158,159],'C09','none')
put([159],'C09','concrete','Можно заметить подачу бетона, но стадию устройства пути этим не установить.')
put([180],'C06','network')
put(range(243,254),'C06','network')
put([251],'C06','none','Прокладка кабеля может выполняться вручную в готовом канале.')
put([252],'C06','concrete_cycle','Составная строка: бетонное основание и монтаж ДГУ. Требуется уточнение подоперации.')
put(range(254,352),'C07','none')
put([263],'C06','earth_haul','В источнике находится под внутренними сетями, но земляные работы отнесены по смыслу к наружным сетям. Проверьте привязку по проекту.')
put([264],'C06','lift','В источнике находится под внутренними сетями. Принят монтаж сборных колодцев; для монолитных измените правило.')
put([265],'C06','network','В источнике находится под внутренними сетями. Проверьте, что речь идёт о наружных трубопроводах.')
put(range(352,381),'C08','none')
put([352,353,354,375],'C08','road')
put([356,378],'C08','concrete')
put([359],'C08','concrete_cycle')
put([368,369,370],'C08','earth')
put([376],'C06','network')
put([377],'C06','none','Кабель может прокладываться вручную.')
put([379],'C08','lift')
put([380],'C07','none')

assert len(raw)==377
assert set(mapping)=={r['source_row'] for r in raw}
code_names = {r['source_code']:r['name'] for r in raw if r['source_code']}
last_code = ''
for i,r in enumerate(raw):
    cid,rid,note=mapping[r['source_row']]
    r.update(class_id=cid,rule_id=rid,mapping_note=note)
    code=r['source_code']
    if code:
        parts=code.split('.')
        parent='.'.join(parts[:-1])
        last_code=code
    else:
        parent=last_code
    ancestors=[]
    parts=parent.split('.') if parent else []
    for j in range(1,len(parts)+1):
        key='.'.join(parts[:j])
        if key in code_names:
            ancestors.append(code_names[key])
    r['path']=' / '.join(ancestors+[r['name']])
    r['parent_code']=parent
    next_code=raw[i+1]['source_code'] if i+1<len(raw) else None
    r['is_group']=bool(code and next_code is not None and (next_code=='' or next_code.startswith(code+'.')))
    r['earlier_profiles']=['E','P'] if cid in ('C04','C05') and rid in ('concrete','concrete_cycle','lift') else []
    r['phase_hint']='Подземная часть' if 46<=r['source_row']<=95 else ('Надземная часть' if 96<=r['source_row']<=241 else '')

# Preserve the hierarchy; section-level selection never inherits a leaf's mandatory machines.
for r in raw:
    if r['is_group']:
        prefix=r['source_code']+'.'
        descendants=[x for x in raw if x['source_row']>r['source_row'] and (x['source_code'].startswith(prefix) or x['parent_code']==r['source_code'] or x['parent_code'].startswith(prefix))]
        r['children']=[x['work_id'] for x in descendants]
        r['group_profiles']=sorted({p for x in descendants for p in rules[x['rule_id']]['profiles']})
        r['group_classes']=sorted({x['class_id'] for x in descendants if x['class_id']!='G00'})
        r['has_unobservable_children']=any(not rules[x['rule_id']]['profiles'] for x in descendants)

sources=[
 {'id':'S1','authors':'Roberts, D.; Golparvar-Fard, M.','year':2019,'title':'End-to-end vision-based detection, tracking and activity analysis of earthmoving equipment filmed at ground level','url':'https://doi.org/10.1016/j.autcon.2019.04.006','read_url':'https://experts.illinois.edu/en/publications/end-to-end-vision-based-detection-tracking-and-activity-analysis-/','basis':'Авторская университетская запись и аннотация','finding':'Работа связывает детекцию и траектории экскаваторов и самосвалов с распознаванием действий через HMM.','use':'Наличие машины и выполнение действия разделены; для развития нужны последовательности кадров.'},
 {'id':'S2','authors':'Zheng, Y.; Seppänen, O.; Masood, M. K.; Törmä, S.','year':2024,'title':'Ontology-Based Construction Process Library for Process States Inference','url':'https://doi.org/10.1007/978-3-031-35399-4_32','read_url':'https://research.aalto.fi/en/publications/ontology-based-construction-process-library-for-process-states-in/','basis':'Университетская запись и аннотация; год по Aalto','finding':'Описана библиотека правил вывода состояния строительных процессов с использованием онтологий и SHACL.','use':'Правила и каталог хранятся отдельно от алгоритма; Excel позволяет проверить каждое сопоставление.'},
 {'id':'S3','authors':'Pal, A.; Lin, J. J.; Hsieh, S. H.; Golparvar-Fard, M.','year':2023,'title':'Automated vision-based construction progress monitoring in built environment through digital twin','url':'https://doi.org/10.1016/j.dibe.2023.100247','read_url':'https://experts.illinois.edu/en/publications/automated-vision-based-construction-progress-monitoring-in-built-/','basis':'Университетская запись и аннотация обзора','finding':'Обзор связывает визуальный мониторинг с планированием, BIM и цифровым двойником строительства.','use':'Состав техники проверяет совместимость; календарное отставание требует данных о плане и фактическом прогрессе.'},
 {'id':'S4','authors':'Sherafat, B.; Rashidi, A.; Lee, Y.-C.; Ahn, C. R.','year':2019,'title':'Automated Activity Recognition of Construction Equipment Using a Data Fusion Approach','url':'https://arxiv.org/abs/1906.02070','read_url':'https://arxiv.org/abs/1906.02070','basis':'Авторский препринт, аннотация','finding':'Распознавание действий объединяет акустические и кинематические данные.','use':'Статус работы и простоя требует дополнительных наблюдений. Веса данного прототипа не взяты из статьи.'},
]
config={'schema_version':1,'source':{'filename':SRC.name,'sha256':hashlib.sha256(SRC.read_bytes()).hexdigest(),'sheet':ws.title,'object_types':object_types,'row_count':len(raw),'codes_repaired':sum(r['code_repaired'] for r in raw)},'machines':machines,'classes':classes,'profiles':profiles,'rules':rules,'thresholds':{'dominant_score':45.,'dominant_margin':15.,'secondary_score':25.,'unexpected_share':.25,'idle_weight':.15},'sources':sources,'works':raw}
OUT.mkdir(parents=True,exist_ok=True)
(OUT/'rules.json').write_text(json.dumps(config,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
with (OUT/'work_mapping.csv').open('w',encoding='utf-8-sig',newline='') as f:
    fields=['work_id','source_row','source_code','name','path','class_id','rule_id','is_group','code_repaired','mapping_note']
    writer=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore')
    writer.writeheader(); writer.writerows(raw)
print(json.dumps({'rows':len(raw),'classes':{k:sum(r['class_id']==k for r in raw) for k in ['G00',*classes]},'groups':sum(r['is_group'] for r in raw),'repaired_codes':config['source']['codes_repaired']},ensure_ascii=False))
